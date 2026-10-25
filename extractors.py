"""
Text extraction and normalization module for SmartLex-S.
Supports TXT, PDF, and DOCX formats with compiler-inspired tokenization.
"""

import os
import re
from typing import List, Tuple, Dict, Set


# Common English stopwords to eliminate non-informative terms in semantic analysis
STOPWORDS: Set[str] = {
    "a", "about", "above", "after", "again", "against", "all", "am", "an", "and", "any", "are", 
    "as", "at", "be", "because", "been", "before", "being", "below", "between", "both", "but", 
    "by", "can", "could", "did", "do", "does", "doing", "down", "during", "each", "few", "for", 
    "from", "further", "had", "has", "have", "having", "he", "her", "here", "hers", "herself", 
    "him", "himself", "his", "how", "i", "if", "in", "into", "is", "it", "its", "itself", "just", 
    "me", "more", "most", "my", "myself", "no", "nor", "not", "now", "of", "off", "on", "once", 
    "only", "or", "other", "our", "ours", "ourselves", "out", "over", "own", "s", "same", "she", 
    "should", "so", "some", "such", "than", "that", "the", "their", "theirs", "them", "themselves", 
    "then", "there", "these", "they", "this", "those", "through", "to", "too", "under", "until", 
    "up", "very", "was", "we", "were", "what", "when", "where", "which", "while", "who", "whom", 
    "why", "will", "with", "would", "you", "your", "yours", "yourself", "yourselves"
}


class ExtractionError(Exception):
    """Custom exception raised when document extraction fails."""
    pass


def extract_text_from_file(filepath: str) -> Tuple[str, str]:
    """
    Extract raw text from a document based on its extension.
    Returns (raw_text, file_type).
    Raises ExtractionError on failure.
    """
    if not os.path.exists(filepath):
        raise ExtractionError(f"File not found: {filepath}")

    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".txt":
        return _extract_txt(filepath), "TXT"
    elif ext == ".pdf":
        return _extract_pdf(filepath), "PDF"
    elif ext == ".docx":
        return _extract_docx(filepath), "DOCX"
    else:
        raise ExtractionError(f"Unsupported file format '{ext}'. Only .txt, .pdf, and .docx are supported.")


def _extract_txt(filepath: str) -> str:
    """Extract text from a plain text file with multiple encoding fallbacks."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    for enc in encodings:
        try:
            with open(filepath, "r", encoding=enc) as f:
                content = f.read()
                return content
        except UnicodeDecodeError:
            continue
        except Exception as e:
            raise ExtractionError(f"Error reading TXT file: {str(e)}")
    raise ExtractionError("Failed to decode text file with standard encodings.")


def _extract_pdf(filepath: str) -> str:
    """Extract text from a PDF file using pypdf."""
    try:
        import pypdf
    except ImportError:
        raise ExtractionError("Missing dependency: 'pypdf'. Run 'pip install pypdf' to enable PDF support.")

    try:
        reader = pypdf.PdfReader(filepath)
        text_parts = []
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        extracted = "\n".join(text_parts).strip()
        if not extracted:
            raise ExtractionError("PDF was opened successfully but contains no extractable text (it might be scanned).")
        return extracted
    except Exception as e:
        if isinstance(e, ExtractionError):
            raise
        raise ExtractionError(f"Error extracting PDF: {str(e)}")


def _extract_docx(filepath: str) -> str:
    """Extract text from a DOCX file using python-docx."""
    try:
        import docx
    except ImportError:
        raise ExtractionError("Missing dependency: 'python-docx'. Run 'pip install python-docx' to enable DOCX support.")

    try:
        doc = docx.Document(filepath)
        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text:
                text_parts.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))
        extracted = "\n".join(text_parts).strip()
        if not extracted:
            raise ExtractionError("DOCX file was opened successfully but contains no readable text.")
        return extracted
    except Exception as e:
        if isinstance(e, ExtractionError):
            raise
        raise ExtractionError(f"Error extracting DOCX: {str(e)}")


def split_into_sentences(text: str) -> List[str]:
    """
    Split text into coherent sentence units.
    Useful for local context, best sentence coverage (C), and snippet evidence.
    """
    if not text:
        return []
    # Normalize line breaks and spaces
    clean = re.sub(r"[\r\n]+", " ", text)
    # Split on sentence terminals followed by space or end
    raw_sentences = re.split(r"(?<=[.!?])\s+", clean)
    sentences = [s.strip() for s in raw_sentences if len(s.strip()) > 5]
    return sentences if sentences else [clean.strip()]


def tokenize_text(text: str, remove_stopwords: bool = True) -> List[str]:
    """
    Compiler-inspired lexical tokenization:
    Extracts identifiers and words (alphanumeric sequences), lowercases them,
    and optionally filters stopwords.
    """
    if not text:
        return []
    # Identify word/identifier tokens
    raw_tokens = re.findall(r"\b[a-zA-Z0-9_]{2,}\b", text.lower())
    if remove_stopwords:
        filtered = [t for t in raw_tokens if t not in STOPWORDS]
        # If removing stopwords would leave an empty list (e.g. query is "what is it"),
        # keep the original tokens so queries are never completely empty
        return filtered if filtered else raw_tokens
    return raw_tokens


def count_token_frequencies(tokens: List[str]) -> Dict[str, int]:
    """Count token occurrences for symbol table / bag-of-words representation."""
    freq: Dict[str, int] = {}
    for t in tokens:
        freq[t] = freq.get(t, 0) + 1
    return freq
